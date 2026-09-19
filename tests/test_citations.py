import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document

from citation_formats import format_citation, normalize_style
from server import app


ARTICLE = {
    "id": "12",
    "title": "عنوان پژوهش 1402 English",
    "venue": "Journal of Testing",
    "year": "1402",
    "type": "مقاله ژورنالی",
    "url": "https://civilica.com/doc/12/",
}


class CitationFormatTests(unittest.TestCase):
    def test_style_aliases_and_formats(self):
        self.assertEqual(normalize_style("APA 7th"), "apa7")
        self.assertIn("(۱۴۰۲)", format_citation(ARTICLE, 1, "apa7", "پژوهشگر نمونه"))
        self.assertIn("(۱۴۰۲)", format_citation({**ARTICLE, "year": "۱۴۰۲"}, 1, "apa7", "پژوهشگر نمونه"))
        self.assertTrue(format_citation(ARTICLE, 2, "vancouver", "پژوهشگر نمونه").startswith("۲."))
        self.assertIn("@misc", format_citation(ARTICLE, 1, "bibtex", "پژوهشگر نمونه"))

    def test_isolating_target_author_preserves_all_authors_by_default(self):
        article = {**ARTICLE, "authors": "رضا خاکپور، ناصر مهردادی، امیر پازوکی"}

        full = format_citation(article, 1, "apa7", "ناصر مهردادی")
        isolated = format_citation(
            article,
            1,
            "apa7",
            "ناصر مهردادی",
            True,
            "ناصر مهردادی",
            True,
        )

        self.assertIn("رضا خاکپور", full)
        self.assertIn("امیر پازوکی", full)
        self.assertEqual(isolated.split(" (")[0], "ناصر مهردادی")


class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()
        self.payload = {
            "profile": {
                "name": "دکتر Soheil",
                "url": "https://civilica.com/p/176225/",
            },
            "articles": [ARTICLE],
            "style": "apa7",
        }

    def test_docx_uses_requested_fonts_rtl_and_digit_rules(self):
        response = self.client.post(
            "/api/export-word",
            json={**self.payload, "file_type": "docx"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.content_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        document = Document(BytesIO(response.data))
        citation = document.paragraphs[-1]
        self.assertIn("۱۴۰۲", citation.text)
        self.assertNotIn("1402", citation.text)
        self.assertIn("https://civilica.com/doc/12/", citation.text)
        self.assertIn("w:bidi", citation._p.xml)
        self.assertIn("B Nazanin", {run.font.name for run in citation.runs})
        self.assertIn("Times New Roman", {run.font.name for run in citation.runs})
        self.assertIn(12.0, {run.font.size.pt for run in citation.runs if run.font.size})
        self.assertIn(11.0, {run.font.size.pt for run in citation.runs if run.font.size})
        self.assertTrue(all('w:val="right"' in paragraph._p.xml for paragraph in document.paragraphs))

    def test_persian_digits_keep_persian_font_runs_and_urls_stay_english(self):
        article = {
            **ARTICLE,
            "title": "عنوان پژوهش ۱۴۰۲ English",
            "year": "۱۴۰۲",
        }
        response = self.client.post(
            "/api/export-word",
            json={**self.payload, "articles": [article], "file_type": "docx"},
        )

        self.assertEqual(response.status_code, 200)
        document = Document(BytesIO(response.data))
        citation = document.paragraphs[-1]
        self.assertIn("۱۴۰۲", citation.text)
        self.assertNotIn("1402", citation.text)
        numeric_runs = [run for run in citation.runs if "۱۴۰۲" in run.text]
        self.assertTrue(numeric_runs)
        self.assertTrue(all(run.font.name == "B Nazanin" for run in numeric_runs))
        self.assertTrue(all(run.font.size.pt == 12.0 for run in numeric_runs))
        url_runs = [run for run in citation.runs if "https://" in run.text]
        self.assertTrue(url_runs)
        self.assertTrue(all(run.font.name == "Times New Roman" for run in url_runs))
        self.assertTrue(all(run.font.size.pt == 11.0 for run in url_runs))

    def test_docx_can_omit_civilica_links(self):
        response = self.client.post(
            "/api/export-word",
            json={**self.payload, "file_type": "docx", "include_links": False},
        )

        self.assertEqual(response.status_code, 200)
        document = Document(BytesIO(response.data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("https://civilica.com", text)
        self.assertIn("صفحهٔ پژوهشگر سیویلیکا", text)

    def test_docx_can_isolate_the_target_author(self):
        article = {
            **ARTICLE,
            "authors": "رضا خاکپور، ناصر مهردادی، امیر پازوکی",
        }
        response = self.client.post(
            "/api/export-word",
            json={
                **self.payload,
                "articles": [article],
                "target_author": "ناصر مهردادی",
                "isolate_author": True,
                "file_type": "docx",
            },
        )

        self.assertEqual(response.status_code, 200)
        document = Document(BytesIO(response.data))
        citation = document.paragraphs[-1].text
        self.assertIn("ناصر مهردادی", citation)
        self.assertNotIn("رضا خاکپور", citation)
        self.assertNotIn("امیر پازوکی", citation)

    def test_doc_compatibility_export_is_word_readable(self):
        response = self.client.post(
            "/api/export-word",
            json={**self.payload, "file_type": "doc"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content_type, "application/msword")
        self.assertTrue(response.data.startswith(b"\xef\xbb\xbf<!doctype html"))
        self.assertIn(b"<head>", response.data)
        self.assertIn(b"text-align: right", response.data)
        self.assertIn("B Nazanin".encode(), response.data)
        self.assertIn("Times New Roman".encode(), response.data)
        self.assertIn("۱۴۰۲".encode(), response.data)
        self.assertIn(b"font-size: 12pt", response.data)
        self.assertIn(b"font-size:11pt", response.data)
        self.assertIn(b"/doc/12/", response.data)


class AuthorEnrichmentEndpointTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_profile_response_does_not_block_on_detail_pages(self):
        parsed = {
            "profile": {
                "id": "176225",
                "name": "پژوهشگر نمونه",
                "url": "https://civilica.com/p/176225/",
            },
            "articles": [{
                "id": "12",
                "title": "مقالهٔ اول",
                "venue": "مجلهٔ نمونه",
                "year": "1402",
                "type": "مقاله ژورنالی",
                "url": "https://civilica.com/doc/12/",
            }],
        }
        with patch("server.fetch_profile_html", return_value="<html></html>"), \
                patch("server.parse_profile_html", return_value=parsed), \
                patch("server.enrich_article_authors") as enrich:
            response = self.client.post(
                "/api/parse-profile",
                json={"url": "https://civilica.com/p/176225/"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["authors_pending"])
        enrich.assert_not_called()

    def test_author_enrichment_returns_authors_for_one_batch(self):
        def fill_authors(articles):
            articles[0]["authors"] = "رضا خاکپور، ناصر مهردادی"

        with patch("server.enrich_article_authors", side_effect=fill_authors):
            response = self.client.post(
                "/api/enrich-authors",
                json={"articles": [{**ARTICLE, "authors": ""}]},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["articles"][0]["authors"],
            "رضا خاکپور، ناصر مهردادی",
        )


if __name__ == "__main__":
    unittest.main()
