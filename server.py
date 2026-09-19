"""Local-first Civilica article extractor with Persian Word exports."""

from __future__ import annotations

import io
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from flask import Flask, jsonify, request, send_file, send_from_directory

from citation_formats import format_citation, normalize_style, style_label
from civilica_parser import parse_article_authors_html, parse_profile_html


ROOT = Path(__file__).resolve().parent
WEB_DIR = ROOT / "web"
MAX_HTML_BYTES = 8 * 1024 * 1024
MAX_ARTICLES = 1000
ARTICLE_AUTHORS_WORKERS = 8
ARTICLE_AUTHORS_TIMEOUT = 15
ARTICLE_DETAIL_MAX_BYTES = 2 * 1024 * 1024
# Keep each proxy request short. The profile endpoint returns the article list
# immediately; the browser fills in co-authors through these small batches.
AUTHOR_ENRICH_BATCH_SIZE = 24
USER_AGENT = "CivilicaPaperExtractor/0.1 (local research utility)"
PERSIAN_FONT = "B Nazanin"
ENGLISH_FONT = "Times New Roman"
WORD_PERSIAN_SIZE = 12
WORD_ENGLISH_SIZE = 11
_WESTERN_DIGITS = "0123456789"
_PERSIAN_DIGITS = "۰۱۲۳۴۵۶۷۸۹"
_ARABIC_DIGITS = "٠١٢٣٤٥٦٧٨٩"
_URL_PATTERN = re.compile(r"https?://[^\s]+", re.IGNORECASE)
_PERSIAN_CHAR_PATTERN = re.compile(r"[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff\ufb50-\ufdff\ufe70-\ufeff]")
_LATIN_CHAR_PATTERN = re.compile(r"[A-Za-z]")
_CIVILICA_HOSTS = {"civilica.com", "www.civilica.com"}

app = Flask(__name__, static_folder=str(WEB_DIR), static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = MAX_HTML_BYTES + 256 * 1024


def allowed_origins() -> set[str]:
    configured = os.environ.get("CIVILICA_ALLOWED_ORIGINS", "")
    origins = {
        origin.strip()
        for origin in configured.split(",")
        if origin.strip()
    }
    origins.update({
        "http://127.0.0.1:5000",
        "http://localhost:5000",
    })
    return origins


@app.after_request
def add_cors_headers(response):
    origin = request.headers.get("Origin", "")
    if origin in allowed_origins():
        response.headers["Access-Control-Allow-Origin"] = origin
        response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
        response.headers["Access-Control-Allow-Headers"] = "Content-Type"
        response.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
        response.headers["Vary"] = "Origin"
    return response


def json_error(message: str, status: int = 400):
    return jsonify({"ok": False, "error": message}), status


def canonical_profile_url(value: str) -> str:
    raw = (value or "").strip()
    parsed = urlparse(raw)
    host = (parsed.hostname or "").lower()
    if parsed.scheme not in {"http", "https"} or host not in {"civilica.com", "www.civilica.com"}:
        raise ValueError("فقط لینک صفحهٔ پژوهشگر سیویلیکا پذیرفته می‌شود.")

    match = re.fullmatch(r"/p/(\d+)/?", parsed.path)
    if not match:
        raise ValueError("لینک باید شبیه https://civilica.com/p/xxxxxx/ باشد.")

    return f"https://civilica.com/p/{match.group(1)}/"


def decode_html(raw: bytes, content_type: str = "") -> str:
    charset_match = re.search(r"charset=([\\w-]+)", content_type or "", re.IGNORECASE)
    encoding = charset_match.group(1) if charset_match else "utf-8"
    return raw.decode(encoding, errors="replace")


def fetch_civilica_html(
    source_url: str,
    *,
    timeout: int = 30,
    max_bytes: int = MAX_HTML_BYTES,
) -> str:
    request_headers = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.5",
        "Accept-Encoding": "identity",
    }
    try:
        with urlopen(Request(source_url, headers=request_headers), timeout=timeout) as response:
            final_host = (urlparse(response.geturl()).hostname or "").lower()
            if final_host not in _CIVILICA_HOSTS:
                raise RuntimeError("سیویلیکا به یک مقصد ناشناس هدایت کرد.")
            raw = response.read(max_bytes + 1)
            if len(raw) > max_bytes:
                raise ValueError("حجم صفحه بیشتر از حد مجاز است.")
            return decode_html(raw, response.headers.get("Content-Type", ""))
    except HTTPError as error:
        if error.code in {401, 403, 429}:
            raise RuntimeError(
                "سیویلیکا دسترسی مستقیم این سرور را محدود کرده است. "
                "حالت «چسباندن HTML» را امتحان کنید."
            ) from error
        raise RuntimeError(f"سیویلیکا با خطای {error.code} پاسخ داد.") from error
    except (URLError, TimeoutError, OSError) as error:
        raise RuntimeError(
            "اتصال به سیویلیکا برقرار نشد. می‌توانید HTML صفحه را در حالت جایگزین وارد کنید."
        ) from error


def fetch_profile_html(source_url: str) -> str:
    return fetch_civilica_html(source_url)


def fetch_article_authors(article: dict[str, object]) -> str:
    """Fetch one trusted Civilica article page and read its citation authors."""

    article_url = str(article.get("url") or "").strip()
    parsed = urlparse(article_url)
    if (
        parsed.scheme not in {"http", "https"}
        or (parsed.hostname or "").lower() not in _CIVILICA_HOSTS
        or not re.fullmatch(r"/doc/\d+/?", parsed.path)
    ):
        return ""

    html = fetch_civilica_html(
        article_url,
        timeout=ARTICLE_AUTHORS_TIMEOUT,
        max_bytes=ARTICLE_DETAIL_MAX_BYTES,
    )
    return parse_article_authors_html(html)


def enrich_article_authors(articles: list[dict[str, object]]) -> None:
    """Add co-authors without making a failed detail request lose an article."""

    pending = [
        article
        for article in articles
        if article.get("url") and not str(article.get("authors") or "").strip()
    ]
    if not pending:
        return

    with ThreadPoolExecutor(max_workers=ARTICLE_AUTHORS_WORKERS) as executor:
        futures = {
            executor.submit(fetch_article_authors, article): article
            for article in pending
        }
        for future in as_completed(futures):
            article = futures[future]
            try:
                authors = future.result()
            except Exception:
                authors = ""
            if authors:
                article["authors"] = authors


def normalized_payload(result: dict[str, object]) -> dict[str, object]:
    articles = result.get("articles") or []
    articles = list(articles)[:MAX_ARTICLES]
    profile = result.get("profile") or {}
    return {
        "ok": True,
        "profile": profile,
        "articles": articles,
        "count": len(articles),
    }


def safe_text(value: object, limit: int = 12000) -> str:
    return str(value or "").strip()[:limit]


def sanitize_articles(value: object) -> list[dict[str, str]]:
    if not isinstance(value, list):
        raise ValueError("فهرست مقاله‌ها معتبر نیست.")

    cleaned: list[dict[str, str]] = []
    for item in value[:MAX_ARTICLES]:
        if not isinstance(item, dict):
            raise ValueError("یکی از مقاله‌های انتخاب‌شده معتبر نیست.")
        title = safe_text(item.get("title"))
        if not title:
            continue
        cleaned.append(
            {
                "id": safe_text(item.get("id"), 120),
                "title": title,
                "venue": safe_text(item.get("venue")),
                "year": safe_text(item.get("year"), 32),
                "type": safe_text(item.get("type"), 120),
                "authors": safe_text(item.get("authors")),
                "url": safe_text(item.get("url"), 1000),
            }
        )
    if not cleaned:
        raise ValueError("حداقل یک مقالهٔ معتبر انتخاب کنید.")
    return cleaned


def boolean_value(value: object, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if not normalized:
        return default
    return normalized not in {"0", "false", "no", "off"}


@app.get("/")
def index():
    return send_from_directory(WEB_DIR, "index.html")


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "service": "CivilicaPulse"})


@app.post("/api/parse-profile")
def parse_profile():
    body = request.get_json(silent=True) or {}
    try:
        source_url = canonical_profile_url(str(body.get("url", "")))
        html = fetch_profile_html(source_url)
        result = parse_profile_html(html, source_url)
        if not result["articles"]:
            return json_error(
                "مقاله‌ای در صفحه پیدا نشد. لینک را بررسی کنید یا HTML صفحه را در حالت جایگزین بچسبانید.",
                422,
            )
        payload = normalized_payload(result)
        payload["authors_pending"] = True
        return jsonify(payload)
    except ValueError as error:
        return json_error(str(error), 400)
    except RuntimeError as error:
        return json_error(str(error), 502)


@app.post("/api/enrich-authors")
def enrich_authors():
    """Resolve co-authors for one short batch after the profile is visible."""

    body = request.get_json(silent=True) or {}
    try:
        articles = sanitize_articles(body.get("articles"))
        if len(articles) > AUTHOR_ENRICH_BATCH_SIZE:
            return json_error(
                f"تعداد مقاله‌های هر مرحله نباید بیشتر از {AUTHOR_ENRICH_BATCH_SIZE} باشد.",
                400,
            )
        enrich_article_authors(articles)
        return jsonify({
            "ok": True,
            "articles": articles,
            "count": len(articles),
        })
    except ValueError as error:
        return json_error(str(error), 400)


@app.post("/api/parse-html")
def parse_html():
    body = request.get_json(silent=True) or {}
    html = str(body.get("html", ""))
    if not html.strip():
        return json_error("کد HTML صفحه را وارد کنید.", 400)
    if len(html.encode("utf-8")) > MAX_HTML_BYTES:
        return json_error("حجم HTML بیشتر از حد مجاز است.", 413)

    raw_source_url = str(body.get("source_url", "")).strip()
    if raw_source_url:
        try:
            source_url = canonical_profile_url(raw_source_url)
        except ValueError:
            source_url = "https://civilica.com/"
    else:
        source_url = "https://civilica.com/"

    result = parse_profile_html(html, source_url)
    if not result["articles"]:
        return json_error(
            "از این HTML مقاله‌ای پیدا نشد. در مرورگر، گزینهٔ View Page Source را کپی کنید.",
            422,
        )
    return jsonify(normalized_payload(result))


def add_bidi(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph_format = paragraph._p.get_or_add_pPr()
    if paragraph_format.find(qn("w:bidi")) is None:
        paragraph_format.append(OxmlElement("w:bidi"))


def persianize_digits(value: object) -> str:
    text = str(value if value is not None else "")
    translation = str.maketrans(
        _WESTERN_DIGITS + _ARABIC_DIGITS,
        _PERSIAN_DIGITS + _PERSIAN_DIGITS,
    )
    return text.translate(translation)


def _script_chunks(value: str):
    current_script = "fa"
    buffer: list[str] = []

    def flush() -> tuple[str, str] | None:
        if not buffer:
            return None
        chunk = "".join(buffer)
        buffer.clear()
        return chunk, current_script

    for character in value:
        if character in "\r\n":
            pending = flush()
            if pending:
                yield pending
            yield "\n", "fa"
            current_script = "fa"
            continue

        if _PERSIAN_CHAR_PATTERN.search(character) or character in (_PERSIAN_DIGITS + _ARABIC_DIGITS):
            detected_script = "fa"
        elif _LATIN_CHAR_PATTERN.search(character):
            detected_script = "en"
        elif character in _WESTERN_DIGITS:
            detected_script = "en"
        else:
            detected_script = current_script

        if buffer and detected_script != current_script:
            pending = flush()
            if pending:
                yield pending
        current_script = detected_script
        buffer.append(character)

    pending = flush()
    if pending:
        yield pending


def word_text_chunks(value: object):
    text = str(value if value is not None else "")
    cursor = 0
    for match in _URL_PATTERN.finditer(text):
        before = text[cursor:match.start()]
        if before:
            for chunk, script in _script_chunks(persianize_digits(before)):
                yield chunk, script, False
        yield match.group(0), "en", True
        cursor = match.end()

    remainder = text[cursor:]
    if remainder:
        for chunk, script in _script_chunks(persianize_digits(remainder)):
            yield chunk, script, False


def set_run_font(
    run,
    size: float | None = None,
    bold: bool | None = None,
    font_name: str = PERSIAN_FONT,
    rtl: bool = True,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    run.font.name = font_name
    for font_slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{font_slot}"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold

    rtl_tag = rpr.find(qn("w:rtl"))
    if rtl and rtl_tag is None:
        from docx.oxml import OxmlElement

        rpr.append(OxmlElement("w:rtl"))
    elif not rtl and rtl_tag is not None:
        rpr.remove(rtl_tag)


def add_word_text(
    paragraph,
    value: object,
    size: float = WORD_PERSIAN_SIZE,
    bold: bool = False,
) -> None:
    for chunk, script, is_url in word_text_chunks(value):
        if chunk == "\n":
            if paragraph.runs:
                paragraph.runs[-1].add_break()
            else:
                paragraph.add_run().add_break()
            continue
        run = paragraph.add_run(chunk)
        if script == "fa":
            set_run_font(run, size=size, bold=bold, font_name=PERSIAN_FONT, rtl=True)
        else:
            set_run_font(
                run,
                size=WORD_ENGLISH_SIZE if size == WORD_PERSIAN_SIZE else max(size - 1, 1),
                bold=bold,
                font_name=ENGLISH_FONT,
                rtl=False,
            )


def build_docx(payload: dict[str, object]) -> io.BytesIO:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = PERSIAN_FONT
    normal_rpr = normal._element.get_or_add_rPr()
    normal_rfonts = normal_rpr.rFonts
    if normal_rfonts is None:
        from docx.oxml import OxmlElement

        normal_rfonts = OxmlElement("w:rFonts")
        normal_rpr.insert(0, normal_rfonts)
    from docx.oxml.ns import qn

    for font_slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_rfonts.set(qn(f"w:{font_slot}"), PERSIAN_FONT)
    normal.font.size = Pt(WORD_PERSIAN_SIZE)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    profile = payload.get("profile") or {}
    profile_name = str(profile.get("name") or "پژوهشگر سیویلیکا")
    source_url = str(profile.get("url") or "")
    articles = list(payload.get("articles") or [])[:MAX_ARTICLES]
    selected_style = normalize_style(payload.get("style"))
    selected_style_label = style_label(selected_style)
    include_links = boolean_value(payload.get("include_links"), default=True)
    target_author = safe_text(payload.get("target_author"), 500)
    isolate_author = boolean_value(payload.get("isolate_author"), default=False)
    citations = [
        format_citation(
            article,
            index,
            selected_style,
            profile_name,
            include_links,
            target_author,
            isolate_author,
        )
        for index, article in enumerate(articles, start=1)
    ]

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_bidi(title)
    add_word_text(title, "فهرست منابع سیویلیکا", size=WORD_PERSIAN_SIZE, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_bidi(subtitle)
    add_word_text(subtitle, profile_name, size=WORD_PERSIAN_SIZE, bold=True)

    info = document.add_paragraph()
    add_bidi(info)
    add_word_text(info, "تعداد مقالات انتخاب‌شده: ", size=WORD_PERSIAN_SIZE, bold=True)
    add_word_text(info, len(articles), size=WORD_PERSIAN_SIZE)
    add_word_text(info, "  |  سبک ارجاع: ", size=WORD_PERSIAN_SIZE, bold=True)
    add_word_text(info, selected_style_label, size=WORD_PERSIAN_SIZE)
    if include_links:
        add_word_text(info, "  |  منبع: ", size=WORD_PERSIAN_SIZE, bold=True)
        add_word_text(info, source_url, size=WORD_PERSIAN_SIZE)
    else:
        add_word_text(info, "  |  منبع: صفحهٔ پژوهشگر سیویلیکا", size=WORD_PERSIAN_SIZE)
    info.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    note = document.add_paragraph()
    add_bidi(note)
    add_word_text(
        note,
        "این فایل بر اساس اطلاعات نمایه‌شده در صفحهٔ پژوهشگر سیویلیکا ساخته شده است. "
        "اگر نام نویسندگان در فهرست عمومی موجود نباشد، نام پژوهشگر پروفایل به‌عنوان "
        "نویسندهٔ جایگزین استفاده می‌شود؛ برای استناد نهایی آن را بررسی کنید.",
        size=WORD_PERSIAN_SIZE,
    )
    note.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for citation in citations:
        citation_paragraph = document.add_paragraph()
        add_bidi(citation_paragraph)
        citation_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        citation_paragraph.paragraph_format.space_after = Pt(7)
        add_word_text(citation_paragraph, citation, size=WORD_PERSIAN_SIZE)

    output = io.BytesIO()
    document.core_properties.title = f"منابع سیویلیکا - {profile_name}"
    document.core_properties.subject = f"فهرست منابع با سبک {selected_style_label}"
    document.save(output)
    output.seek(0)
    return output


def html_word_markup(
    value: object,
    size: float = WORD_PERSIAN_SIZE,
    bold: bool = False,
) -> str:
    from html import escape

    markup: list[str] = []
    for chunk, script, is_url in word_text_chunks(value):
        if chunk == "\n":
            markup.append("<br>")
            continue
        font_name = PERSIAN_FONT if script == "fa" else ENGLISH_FONT
        font_size = (
            size
            if script == "fa"
            else WORD_ENGLISH_SIZE
            if size == WORD_PERSIAN_SIZE
            else max(size - 1, 1)
        )
        direction = "rtl" if script == "fa" and not is_url else "ltr"
        weight = "font-weight:700;" if bold else ""
        markup.append(
            f'<span dir="{direction}" style="font-family:\'{escape(font_name)}\';'
            f'font-size:{font_size}pt;direction:{direction};unicode-bidi:embed;'
            f'mso-ansi-font-family:\'{escape(font_name)}\';'
            f'mso-fareast-font-family:\'{escape(font_name)}\';'
            f'mso-bidi-font-family:\'{escape(font_name)}\';{weight}">'
            f'{escape(chunk)}</span>'
        )
    return "".join(markup)


def build_word_html(payload: dict[str, object]) -> io.BytesIO:
    from html import escape

    profile = payload.get("profile") or {}
    profile_name = str(profile.get("name") or "پژوهشگر سیویلیکا")
    source_url = str(profile.get("url") or "")
    articles = list(payload.get("articles") or [])[:MAX_ARTICLES]
    selected_style = normalize_style(payload.get("style"))
    selected_style_label = style_label(selected_style)
    include_links = boolean_value(payload.get("include_links"), default=True)
    target_author = safe_text(payload.get("target_author"), 500)
    isolate_author = boolean_value(payload.get("isolate_author"), default=False)
    citations = [
        format_citation(
            article,
            index,
            selected_style,
            profile_name,
            include_links,
            target_author,
            isolate_author,
        )
        for index, article in enumerate(articles, start=1)
    ]
    citation_markup = "".join(
        '<p class="citation" dir="rtl" align="right">'
        + html_word_markup(citation, size=WORD_PERSIAN_SIZE)
        + "</p>"
        for citation in citations
    )
    title_markup = html_word_markup(
        "فهرست منابع سیویلیکا", size=WORD_PERSIAN_SIZE, bold=True
    )
    profile_markup = html_word_markup(profile_name, size=WORD_PERSIAN_SIZE, bold=True)
    meta_markup = (
        html_word_markup("سبک ارجاع: ", size=WORD_PERSIAN_SIZE, bold=True)
        + html_word_markup(selected_style_label, size=WORD_PERSIAN_SIZE)
        + html_word_markup(" | تعداد: ", size=WORD_PERSIAN_SIZE, bold=True)
        + html_word_markup(len(articles), size=WORD_PERSIAN_SIZE)
        + html_word_markup(" | منبع: ", size=WORD_PERSIAN_SIZE, bold=True)
        + (
            html_word_markup(source_url, size=WORD_PERSIAN_SIZE)
            if include_links
            else html_word_markup("صفحهٔ پژوهشگر سیویلیکا", size=WORD_PERSIAN_SIZE)
        )
    )
    note_markup = html_word_markup(
        "این فایل بر اساس اطلاعات نمایه‌شده در صفحهٔ پژوهشگر سیویلیکا ساخته شده است. "
        "اگر نام نویسندگان در فهرست عمومی موجود نباشد، نام پژوهشگر پروفایل به‌عنوان "
        "نویسندهٔ جایگزین استفاده می‌شود؛ برای استناد نهایی آن را بررسی کنید.",
        size=WORD_PERSIAN_SIZE,
    )
    html = f"""<!doctype html>
<html lang="fa" dir="rtl">
<head>
<meta charset="utf-8">
<title>منابع سیویلیکا - {escape(profile_name)}</title>
<style>
  @page {{ margin: 2cm; }}
  html, body {{ direction: rtl; }}
  body {{ font-family: "B Nazanin", Tahoma, Arial, sans-serif; mso-bidi-font-family: "B Nazanin"; direction: rtl; text-align: right; font-size: 12pt; line-height: 1.7; }}
  h1 {{ margin: 0 0 8pt 0; direction: rtl; text-align: right; font-family: "B Nazanin"; font-size: 12pt; }}
  p {{ direction: rtl; text-align: right; font-family: "B Nazanin"; font-size: 12pt; }}
  .meta {{ color: #475569; font-size: 12pt; }}
  .citation {{ margin: 0 0 10pt 0; padding: 0; text-indent: 0; direction: rtl; text-align: right !important; }}
</style>
</head>
<body dir="rtl" align="right">
<h1 dir="rtl" align="right">{title_markup}</h1>
<p dir="rtl" align="right">{profile_markup}</p>
<p class="meta" dir="rtl" align="right">{meta_markup}</p>
<p class="meta" dir="rtl" align="right">{note_markup}</p>
{citation_markup}
</body>
</html>"""
    output = io.BytesIO(("\ufeff" + html).encode("utf-8"))
    output.seek(0)
    return output


@app.post("/api/export-word")
@app.post("/api/export-docx")
def export_word():
    body = request.get_json(silent=True) or {}
    articles = body.get("articles")
    if not isinstance(articles, list) or not articles:
        return json_error("حداقل یک مقاله را انتخاب کنید.", 400)
    if len(articles) > MAX_ARTICLES:
        return json_error("تعداد مقاله‌های انتخاب‌شده بیش از حد مجاز است.", 400)

    profile = body.get("profile") if isinstance(body.get("profile"), dict) else {}
    try:
        sanitized_payload = {
            "profile": {
                "name": safe_text(profile.get("name"), 500),
                "url": safe_text(profile.get("url"), 1000),
            },
            "articles": sanitize_articles(articles),
            "style": normalize_style(body.get("style")),
            "include_links": boolean_value(body.get("include_links"), default=True),
            "target_author": safe_text(body.get("target_author"), 500),
            "isolate_author": boolean_value(body.get("isolate_author"), default=False),
        }
    except ValueError as error:
        return json_error(str(error), 400)

    file_type = "doc" if str(body.get("file_type", "")).lower() == "doc" else "docx"
    try:
        if file_type == "doc":
            output = build_word_html(sanitized_payload)
            mimetype = "application/msword"
        else:
            output = build_docx(sanitized_payload)
            mimetype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    except ValueError as error:
        return json_error(str(error), 400)
    except ImportError:
        return json_error(
            "کتابخانهٔ ساخت فایل ورد نصب نیست. requirements.txt را نصب کنید.",
            503,
        )
    except Exception:
        app.logger.exception("Word export failed")
        return json_error("ساخت فایل ورد انجام نشد.", 500)

    filename = f"civilica-references-{normalize_style(body.get('style'))}-{datetime.now(timezone.utc):%Y%m%d}.{file_type}"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=mimetype,
    )


@app.errorhandler(413)
def request_too_large(_error):
    return json_error("حجم درخواست بیشتر از حد مجاز است.", 413)


if __name__ == "__main__":
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "5000"))
    app.run(host=host, port=port, debug=False)
